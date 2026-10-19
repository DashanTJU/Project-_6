

from __future__ import annotations

import argparse
import copy
import dataclasses
import json
import logging
import os
import platform
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd

import config as C
import data_loader
import evaluation as ev
import fusion
import models as model_lib
import preprocessing as prep
import statistical_tests as st
import synthetic_data
import training
import visualization as viz
import window_generator as wg
from config import (HYBRID_MODELS, STANDALONE_MODELS, TARGET_FOLDERS, TARGETS, AblationConfig, ConfigError,
                    ExperimentConfig, PlantConfig)

STAGES = ("preprocess", "train", "evaluate", "plot")
FUSION_METHODS = ("UEBF", "SCAF", "DALF")
SMOKE_DAYS = 400



class Paths:
    """Output layout: outputs/{logs,data,models,predictions,tables,figures}/<plant>/<target>/..."""

    def __init__(self, root: str | Path):
        self.root = Path(root)

    @staticmethod
    def _mk(path: Path) -> Path:
        path.mkdir(parents=True, exist_ok=True)
        return path

    @property
    def logs(self) -> Path:
        return self._mk(self.root / "logs")

    @property
    def tables(self) -> Path:
        return self._mk(self.root / "tables")

    @property
    def figures(self) -> Path:
        return self._mk(self.root / "figures")

    def data(self, plant: str, target: str) -> Path:
        return self._mk(self.root / "data" / plant / TARGET_FOLDERS[target])

    def tables_for(self, plant: str, target: str) -> Path:
        return self._mk(self.root / "tables" / plant / TARGET_FOLDERS[target])

    def figures_for(self, plant: str, target: str) -> Path:
        return self._mk(self.root / "figures" / plant / TARGET_FOLDERS[target])

    def predictions_for(self, plant: str, target: str) -> Path:
        return self._mk(self.root / "predictions" / plant / TARGET_FOLDERS[target])

    def run_dirs(self, plant: str, target: str) -> dict[str, Path]:
        base = self.root / "models" / plant / TARGET_FOLDERS[target]
        return {"models": self._mk(base / "checkpoints"), "histories": self._mk(base / "histories"),
                "summaries": self._mk(base / "summaries"),
                "predictions": self._mk(self.predictions_for(plant, target) / "runs")}



def _json_default(obj):
    if isinstance(obj, np.integer):
        return int(obj)
    if isinstance(obj, np.floating):
        return float(obj) if np.isfinite(obj) else None
    if isinstance(obj, np.bool_):
        return bool(obj)
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    if isinstance(obj, (pd.Timestamp, np.datetime64)):
        return str(pd.Timestamp(obj))
    return str(obj)


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, default=_json_default, ensure_ascii=False), encoding="utf-8")


def save_table(df: pd.DataFrame, path: Path) -> pd.DataFrame:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    return df


def write_excel(tables: dict[str, pd.DataFrame], path: Path, log: logging.Logger) -> None:
    """One workbook, one sheet per table (sheet names cut to Excel's 31 characters)."""
    try:
        import openpyxl  
    except ImportError:
        log.warning("openpyxl is not installed: %s skipped (the CSV files were written).", path.name)
        return
    used: set[str] = set()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in tables.items():
            sheet, i = name[:31], 1
            while sheet in used:
                suffix = f"_{i}"
                sheet, i = name[:31 - len(suffix)] + suffix, i + 1
            used.add(sheet)
            df.to_excel(writer, sheet_name=sheet, index=False)
    log.info("Excel workbook written: %s", path)


def setup_logging(log_dir: Path) -> logging.Logger:
    logger = logging.getLogger("pv_dalf")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s %(levelname)-7s %(message)s", "%Y-%m-%d %H:%M:%S")
    for handler in (logging.StreamHandler(sys.stdout), logging.FileHandler(log_dir / "run.log", encoding="utf-8")):
        handler.setFormatter(formatter)
        logger.addHandler(handler)
    return logger


def environment_info() -> dict:
    """Hardware and software of the execution environment (measured, nothing assumed)."""
    import importlib.metadata as metadata
    versions = {}
    for package in ("numpy", "pandas", "scikit-learn", "scipy", "matplotlib", "tensorflow", "keras", "joblib",
                    "openpyxl", "psutil"):
        try:
            versions[package] = metadata.version(package)
        except metadata.PackageNotFoundError:
            versions[package] = None
    info = {"python": platform.python_version(), "os": platform.platform(), "machine": platform.machine(),
            "processor": platform.processor() or None, "logical_cpus": os.cpu_count(), "versions": versions}
    try:
        with open("/proc/cpuinfo", encoding="utf-8") as handle:
            names = [line.split(":", 1)[1].strip() for line in handle if line.startswith("model name")]
        info["cpu_model"] = names[0] if names else None
    except OSError:
        info["cpu_model"] = None
    try:
        import psutil
        info["ram_total_gb"] = round(psutil.virtual_memory().total / 1e9, 2)
        info["physical_cpus"] = psutil.cpu_count(logical=False)
    except ImportError:
        info["ram_total_gb"] = None
    return info


def peak_memory_mb() -> float | None:
    """Peak resident memory of this process so far (MB), where the platform reports it."""
    try:
        import resource
        peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        return round(peak / (1024.0 * 1024.0) if sys.platform == "darwin" else peak / 1024.0, 1)
    except ImportError:
        try:
            import psutil
            info = psutil.Process().memory_info()
            return round(getattr(info, "peak_wset", info.rss) / 1e6, 1)
        except ImportError:
            return None


def record_stage(paths: Paths, stage: str, seconds: float) -> None:
    """Keep the latest runtime of each stage and a pipeline total."""
    file = paths.logs / "stage_runtime.json"
    data = json.loads(file.read_text(encoding="utf-8")) if file.exists() else {}
    data[stage] = {"seconds": round(seconds, 3), "finished": pd.Timestamp.now().isoformat(timespec="seconds"),
                   "peak_process_memory_mb": peak_memory_mb()}
    write_json(file, data)
    rows = [{"stage": key, **value} for key, value in data.items()]
    rows.append({"stage": "total (latest run of each stage)", "seconds": round(sum(v["seconds"] for v in data.values()), 3)})
    save_table(pd.DataFrame(rows), paths.tables / "pipeline_runtime.csv")


def energy_function(cfg: ExperimentConfig):
    def to_daily_energy(blocks):
        return ev.aggregate_daily_energy(blocks, cfg.data.power_measurement, cfg.window.horizon_hours)
    return to_daily_energy


def family_of(model: str) -> str:
    if model in FUSION_METHODS:
        return "fusion"
    return model_lib.model_family(model)


# configuration
def parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="PV power and solar-irradiance forecasting with DALF.")
    parser.add_argument("--stage", choices=STAGES + ("all",), default="all", help="pipeline stage (default: all)")
    parser.add_argument("--config", help="JSON configuration file (start from config_example.json)")
    parser.add_argument("--write-default-config", metavar="PATH", help="write the full default configuration and exit")
    parser.add_argument("--smoke-test", action="store_true", help="synthetic data, small models, 2 seeds, 2 epochs")
    parser.add_argument("--plants", nargs="+", help="subset of configured plant names")
    parser.add_argument("--targets", nargs="+", choices=TARGETS, help="power and/or irradiance")
    parser.add_argument("--models", nargs="+", choices=STANDALONE_MODELS + HYBRID_MODELS, help="models to train/evaluate")
    parser.add_argument("--seeds", nargs="+", type=int, help="random seeds")
    parser.add_argument("--epochs", type=int, help="maximum epochs")
    parser.add_argument("--input-hours", type=int, help="history length (default 168)")
    parser.add_argument("--horizon-hours", type=int, help="forecast horizon (default 24)")
    parser.add_argument("--window-days", type=int, help="DALF rolling window in completed days (default 30)")
    parser.add_argument("--output-dir", help="output directory")
    parser.add_argument("--resume", action="store_true", help="reuse finished training runs with identical settings")
    parser.add_argument("--confirm-site-settings", action="store_true",
                        help="confirm units, timestamp convention and physical limits for real data")
    parser.add_argument("--enable-mse-dalf", action="store_true", help="add the separate MSE-based DALF experiment")
    parser.add_argument("--skip-ablation", action="store_true", help="no ablation or sensitivity experiments")
    return parser.parse_args(argv)


def make_smoke_config(cfg: ExperimentConfig, output_dir: str) -> ExperimentConfig:
    """Small, fast settings on SYNTHETIC data (code check only)."""
    cfg = copy.deepcopy(cfg)
    out = Path(output_dir)
    files = synthetic_data.write_synthetic_inputs(out / "synthetic_inputs", days=SMOKE_DAYS)
    site = "SYNTHETIC smoke-test series (not plant data)"
    cfg.plants = [PlantConfig("plant_50mw", 50.0, site, str(files["plant_50mw"])),
                  PlantConfig("plant_100mw", 100.0, site, str(files["plant_100mw"]), sheet_name="hourly")]
    cfg.synthetic_data = True
    cfg.output.output_dir = str(out)
    cfg.window.training_origin_stride_hours = 6
    cfg.model.conv_filters, cfg.model.recurrent_units, cfg.model.dense_units = 4, 8, [8]
    cfg.training.epochs, cfg.training.seeds = 2, [42, 7]
    cfg.training.early_stopping_patience, cfg.training.verbose = 1, 0
    cfg.ablation = AblationConfig(dalf_window_days=[7, 14, 30], mse_dalf=True, interpolation_limits_hours=[0, 1, 3],
                                  input_masking_fractions=[0.05], activations=["relu"], optimizers=["adam"],
                                  variant_seeds=[42])
    return cfg


def build_config(args: argparse.Namespace) -> ExperimentConfig:
    cfg = C.load_config(args.config)
    if args.smoke_test:
        cfg = make_smoke_config(cfg, args.output_dir or "outputs_smoke")
    if args.plants:
        known = {p.name for p in cfg.plants}
        unknown = sorted(set(args.plants) - known)
        if unknown:
            raise ConfigError(f"Unknown plants {unknown}; configured plants: {sorted(known)}")
        cfg.plants = [p for p in cfg.plants if p.name in args.plants]
    if args.targets:
        cfg.targets = list(args.targets)
    if args.models:
        cfg.model.models = list(args.models)
    if args.seeds:
        cfg.training.seeds = list(args.seeds)
    if args.epochs:
        cfg.training.epochs = args.epochs
    if args.input_hours:
        cfg.window.input_hours = args.input_hours
    if args.horizon_hours:
        cfg.window.horizon_hours = args.horizon_hours
    if args.window_days:
        cfg.fusion.window_days = args.window_days
        cfg.fusion.min_history_days = min(cfg.fusion.min_history_days, args.window_days)
    if args.output_dir:
        cfg.output.output_dir = args.output_dir
    if args.confirm_site_settings:
        cfg.site_settings_confirmed = True
    if args.enable_mse_dalf:
        cfg.ablation.mse_dalf = True
    if args.skip_ablation:
        cfg.ablation = AblationConfig(dalf_window_days=[], mse_dalf=False, interpolation_limits_hours=[],
                                      input_masking_fractions=[], activations=[], optimizers=[])
    C.require_valid(cfg)
    return cfg


# preprocess
def stage_preprocess(cfg: ExperimentConfig, paths: Paths, log: logging.Logger, args) -> None:
    quality, counts, sensitivity = [], [], []
    block_start = C.daily_block_start_hour(cfg)
    w = cfg.window
    for plant in cfg.plants:
        raw, source = data_loader.load_table(plant, cfg.data)
        log.info("%s: %d raw rows from %s%s", plant.name, len(raw), source["file"],
                 f" (sheet {source['sheet']})" if source.get("sheet") is not None else "")
        if plant.reference_raw_rows is not None and len(raw) != plant.reference_raw_rows:
            log.warning("%s: %d raw rows differ from the reference %d", plant.name, len(raw), plant.reference_raw_rows)
        for target in cfg.targets:
            folder = paths.data(plant.name, target)
            clean, report = prep.clean_target_series(raw, plant, target, cfg)
            report["source"] = source
            train_end, val_end = prep.chronological_boundaries(len(clean), cfg)
            split = prep.split_report(clean, train_end, val_end, plant, cfg)
            features = prep.variables_for_target(target)
            feature_scaler, target_scaler = prep.fit_scalers(clean.iloc[:train_end], features, target)
            prep.save_scalers(feature_scaler, target_scaler, features, target, train_end, folder / "scalers")
            sets = wg.build_window_sets(clean, features, target, feature_scaler, target_scaler, train_end, val_end,
                                        w.input_hours, w.horizon_hours, w.training_origin_stride_hours, block_start)
            summary = wg.window_summary(sets, train_end, val_end, w.input_hours)
            write_json(folder / "preprocessing_report.json", report)
            write_json(folder / "split_report.json", split)
            write_json(folder / "window_report.json", summary)
            labels = np.array(["train"] * train_end + ["validation"] * (val_end - train_end)
                              + ["test"] * (len(clean) - val_end))
            out = clean.copy()
            out.insert(1, "partition", labels)
            out.to_csv(folder / "cleaned_series.csv", index=False)
            normalized = pd.DataFrame(feature_scaler.transform(clean[features].to_numpy(dtype=float)),
                                      columns=[f"{f}_scaled_input" for f in features])
            normalized.insert(0, "timestamp", clean["timestamp"].to_numpy())
            normalized.insert(1, "partition", labels)
            normalized[f"{target}_scaled_target"] = target_scaler.transform(clean[[target]].to_numpy(dtype=float)).ravel()
            normalized.to_csv(folder / "normalized_series.csv", index=False)
            checks = wg.verify_window_sets(sets, clean, train_end, val_end, w.input_hours, w.horizon_hours, block_start)
            write_json(folder / "window_checks.json", checks)
            wg.save_window_sets(sets, folder / "windows.npz")
            wg.window_index_frame(sets, w.input_hours).to_csv(folder / "window_index.csv", index=False)
            quality.append(prep.quality_summary_row(report))
            row = {"plant": plant.name, "target": target, "cleaned_rows": len(clean)}
            for name, part in split["partitions"].items():
                row.update({f"{name}_rows": part["rows"], f"{name}_index_range": f"{part['index_first']}-{part['index_last']}",
                            f"{name}_first_timestamp": part["first_timestamp"], f"{name}_last_timestamp": part["last_timestamp"],
                            f"{name}_windows": summary[name]["windows"]})
            row["reference_check"] = json.dumps(split["reference_check"]) if isinstance(split["reference_check"], dict) \
                else split["reference_check"]
            counts.append(row)
            log.info("  %s/%s: cleaned %d rows (reference %s); windows train %d, validation %d days, test %d days; %s",
                     plant.name, target, len(clean), plant.reference_cleaned_rows, sets["train"].n,
                     sets["validation"].n, sets["test"].n, row["reference_check"])
            if sets["validation"].n < cfg.fusion.window_days:
                log.warning("  %s/%s: only %d validation days; the first test day cannot use a full %d-day DALF window",
                            plant.name, target, sets["validation"].n, cfg.fusion.window_days)
            for limit in cfg.ablation.interpolation_limits_hours:  # data-level missing-data sensitivity
                c2, r2 = prep.clean_target_series(raw, plant, target, cfg, max_gap_hours=limit)
                a2, b2 = prep.chronological_boundaries(len(c2), cfg)
                fs2, ts2 = prep.fit_scalers(c2.iloc[:a2], features, target)
                s2 = wg.build_window_sets(c2, features, target, fs2, ts2, a2, b2, w.input_hours, w.horizon_hours,
                                          w.training_origin_stride_hours, block_start)
                interp = next(s for s in r2["steps"] if s["step"].startswith("linear"))["variables"]
                sensitivity.append({
                    "plant": plant.name, "target": target, "max_interpolation_gap_hours": limit,
                    "is_main_setting": limit == cfg.cleaning.max_interpolation_gap_hours,
                    "values_interpolated": sum(v["values_interpolated"] for v in interp.values()),
                    "unresolved_hours_removed": r2["unresolved_hours"], "cleaned_rows": len(c2),
                    "train_windows": s2["train"].n, "validation_days": s2["validation"].n, "test_days": s2["test"].n})
    tables = {"data_quality_summary": pd.DataFrame(quality), "sample_counts": pd.DataFrame(counts)}
    if sensitivity:
        tables["ablation_interpolation_limits"] = pd.DataFrame(sensitivity)
    for name, df in tables.items():
        save_table(df, paths.tables / f"{name}.csv")
    write_excel(tables, paths.tables / "preprocessing_tables.xlsx", log)


# train
def training_jobs(cfg: ExperimentConfig) -> list[dict]:
    """Main runs (every model x seed) plus activation/optimizer variants of the hybrid models."""
    masking = tuple(cfg.ablation.input_masking_fractions)
    jobs = [{"model": m, "seed": s, "variant": "main", "model_cfg": cfg.model, "training_cfg": cfg.training,
             "masking": masking} for m in cfg.model.models for s in cfg.training.seeds]
    hybrids = [m for m in cfg.model.models if m in HYBRID_MODELS]
    for activation in cfg.ablation.activations:
        if activation != cfg.model.activation:
            mcfg = dataclasses.replace(cfg.model, activation=activation)
            jobs += [{"model": m, "seed": s, "variant": f"activation-{activation}", "model_cfg": mcfg,
                      "training_cfg": cfg.training, "masking": ()} for m in hybrids for s in cfg.ablation.variant_seeds]
    for optimizer in cfg.ablation.optimizers:
        if optimizer != cfg.training.optimizer:
            tcfg = dataclasses.replace(cfg.training, optimizer=optimizer)
            jobs += [{"model": m, "seed": s, "variant": f"optimizer-{optimizer}", "model_cfg": cfg.model,
                      "training_cfg": tcfg, "masking": ()} for m in hybrids for s in cfg.ablation.variant_seeds]
    return jobs


def stage_train(cfg: ExperimentConfig, paths: Paths, log: logging.Logger, args) -> None:
    C.require_confirmed_site_settings(cfg)
    try:
        determinism = training.configure_determinism(cfg.training.deterministic_ops)
        gpus = [d.name for d in model_lib._tf().config.list_physical_devices("GPU")]
    except ImportError as exc:
        raise RuntimeError("TensorFlow is required for training: pip install -r requirements.txt") from exc
    write_json(paths.logs / "train_environment.json", {**environment_info(), "determinism": determinism, "gpus": gpus})
    log.info("Determinism: %s; GPUs: %s", determinism["note"], gpus or "none")
    jobs = training_jobs(cfg)
    total = len(cfg.plants) * len(cfg.targets) * len(jobs)
    done, trained_seconds, failures = 0, [], []
    for plant in cfg.plants:
        for target in cfg.targets:
            folder = paths.data(plant.name, target)
            if not (folder / "windows.npz").exists():
                raise FileNotFoundError(f"{folder / 'windows.npz'} not found; run --stage preprocess first.")
            sets = wg.load_window_sets(folder / "windows.npz")
            _, target_scaler = prep.load_scalers(folder / "scalers")
            run_dirs = paths.run_dirs(plant.name, target)
            for job in jobs:
                done += 1
                label = f"[{done}/{total}] {plant.name}/{target}/{job['model']}/{job['variant']}/seed {job['seed']}"
                try:
                    result = training.train_single_run(
                        plant=plant.name, target=target, model_type=job["model"], seed=job["seed"], windows=sets,
                        target_scaler=target_scaler, cfg=cfg, run_dirs=run_dirs, variant=job["variant"],
                        model_cfg=job["model_cfg"], training_cfg=job["training_cfg"], masking_fractions=job["masking"],
                        resume=args.resume, logger=log)
                except Exception as exc:  
                    log.error("%s FAILED: %s", label, exc)
                    failures.append({"run": label, "error": f"{type(exc).__name__}: {exc}"})
                    continue
                record = result["record"]
                if not record.get("reused"):
                    trained_seconds.append(record["training_seconds"])
                mean = float(np.mean(trained_seconds)) if trained_seconds else 0.0
                log.info("%s: %s, best epoch %s of %s, %.1f s; projected remaining %.2f h", label,
                         "reused" if record.get("reused") else "trained", record["best_epoch"], record["epochs_run"],
                         record["training_seconds"], (total - done) * mean / 3600)
    write_json(paths.logs / "training_failures.json", failures)
    if failures:
        raise RuntimeError(f"{len(failures)} training runs failed; see {paths.logs / 'training_failures.json'}")


# evaluate
def load_seed_predictions(run_dirs: dict, model: str, variant: str, seeds, sets: dict, masking) -> dict:
    out = {"validation": [], "test": [], "records": [], "masked": {f: [] for f in masking}}
    missing = []
    for seed in seeds:
        name = f"{model}_{variant}_seed{seed}"
        npz, rec = run_dirs["predictions"] / f"{name}.npz", run_dirs["histories"] / f"{name}_record.json"
        if not (npz.exists() and rec.exists()):
            missing.append(name)
            continue
        with np.load(npz) as data:
            for p in ("validation", "test"):
                if not np.array_equal(data[f"{p}_origin_ns"], sets[p].origin_times.astype(np.int64)):
                    raise AssertionError(f"{name}: {p} forecast origins differ from the current windows; retrain "
                                         "(--stage train without --resume)")
                if data[p].shape != sets[p].Y.shape or not np.all(np.isfinite(data[p])):
                    raise AssertionError(f"{name}: {p} predictions must be finite with shape {sets[p].Y.shape}")
                out[p].append(data[p])
            for f in masking:
                if f"test_masked_{f}" in data:
                    out["masked"][f].append(data[f"test_masked_{f}"])
        out["records"].append(json.loads(rec.read_text(encoding="utf-8")))
    if missing:
        raise FileNotFoundError(f"Missing trained runs (run --stage train): {missing}")
    for p in ("validation", "test"):
        out[p] = np.stack(out[p])
    out["masked"] = {f: np.stack(v) for f, v in out["masked"].items() if len(v) == len(seeds)}
    return out


def prediction_frame(plant: str, target: str, ws: wg.WindowSet, forecasts: dict[str, np.ndarray], unit: str,
                     dalf_modes: pd.Series | None = None) -> pd.DataFrame:
    """Hourly prediction table: forecast origin, target timestamp, lead time, actual, forecasts, residuals."""
    n, h = ws.Y.shape
    frame = pd.DataFrame({"plant": plant, "target": target, "partition": ws.partition,
                          "forecast_origin": np.repeat(ws.origin_times, h), "timestamp": ws.target_times().reshape(-1),
                          "lead_hour": np.tile(np.arange(1, h + 1), n), "actual": ws.Y.reshape(-1)})
    for name, values in forecasts.items():
        frame[name] = np.asarray(values, dtype=float).reshape(-1)
    for name in FUSION_METHODS:
        if name in forecasts:
            frame[f"residual_{name}"] = frame["actual"] - frame[name]
    frame["unit"] = unit
    if dalf_modes is not None:
        frame["dalf_window_mode"] = np.repeat(dalf_modes.to_numpy(), h)
    return frame


def evaluate_target(cfg: ExperimentConfig, paths: Paths, plant: PlantConfig, target: str,
                    log: logging.Logger) -> dict:
    folder = paths.data(plant.name, target)
    if not (folder / "windows.npz").exists():
        raise FileNotFoundError(f"{folder / 'windows.npz'} not found; run --stage preprocess first.")
    sets = wg.load_window_sets(folder / "windows.npz")
    _, target_scaler = prep.load_scalers(folder / "scalers")
    run_dirs = paths.run_dirs(plant.name, target)
    val, test = sets["validation"], sets["test"]
    unit, e_unit = C.target_unit(cfg, target), C.energy_unit(cfg)
    resolution = cfg.fusion.resolution[target]
    energy_fn = energy_function(cfg) if target == "power" and e_unit else None
    f = cfg.fusion
    base = list(f.base_models)
    checks: list[dict] = []

    def check(name, ok, detail=""):
        checks.append({"plant": plant.name, "target": target, "check": name, "status": "ok" if ok else "FAILED",
                       "detail": detail})
        if not ok:
            raise AssertionError(f"{plant.name}/{target}: {name} failed. {detail}")

    masking = tuple(cfg.ablation.input_masking_fractions)
    loaded = {m: load_seed_predictions(run_dirs, m, "main", cfg.training.seeds, sets, masking) for m in cfg.model.models}
    check("prediction origins match the 168 h -> 24 h windows", True, "checked for every run")
    ens = {m: {p: loaded[m][p].mean(axis=0) for p in ("validation", "test")} for m in cfg.model.models}

    # fusion
    origins = np.concatenate([val.origin_times, test.origin_times])
    partitions = np.array(["validation"] * val.n + ["test"] * test.n, dtype=object)
    actual_blocks = np.concatenate([val.Y, test.Y])

    def stack(source: dict, parts=("validation", "test")):
        return np.concatenate([np.stack([source[m][p] for m in base], axis=-1) for p in parts])

    base_blocks = stack(ens)

    def run_dalf(blocks, window_days=f.window_days, objective="mae"):
        return fusion.rolling_dalf(actual_blocks, blocks, origins, partitions, window_days=window_days,
                                   resolution=resolution, energy_fn=energy_fn, objective=objective,
                                   window_definition=f.window_definition,
                                   min_history_days=min(f.min_history_days, window_days),
                                   tolerance=f.weight_tolerance, model_names=tuple(base))

    dalf_fc, dalf_w = run_dalf(base_blocks)
    W = dalf_w[base].to_numpy()
    check("DALF weights non-negative and sum to one", bool(np.all(W >= -f.weight_tolerance)
                                                           and np.allclose(W.sum(axis=1), 1.0, atol=f.weight_tolerance)))
    has = dalf_w["window_days_used"] > 0
    check("no actual value of the forecast day (or later) in any DALF window",
          bool((dalf_w.loc[has, "window_last_hour"] < dalf_w.loc[has, "forecast_origin"]).all()))
    first_test = dalf_w[dalf_w["partition"] == "test"].iloc[0]
    if val.n >= f.window_days and f.window_definition == "last_n_completed_days":
        check("first test day uses the final completed validation days",
              first_test["window_days_used"] == f.window_days and first_test["window_partitions"] == "validation")

    scaf_est = fusion.static_weights(val.Y, stack(ens, ("validation",)), resolution, energy_fn, "mae", f.weight_tolerance)
    reported = f.scaf_reported_coefficients.get(plant.name)
    use_reported = f.scaf_mode == "reported" and f.scaf_reported_confirmed and f.scaf_reported_target == target
    if use_reported:
        w_scaf, scaf_source = np.array([reported[m] for m in base], dtype=float), "reported coefficients (confirmed)"
        w_scaf = w_scaf / w_scaf.sum()
    else:
        w_scaf, scaf_source = np.asarray(scaf_est["weights"]), "estimated on the validation period (MAE)"
    scaf_rows = [{"plant": plant.name, "target": target, "coefficients": "estimated on validation", "used": not use_reported,
                  "resolution": resolution, "validation_objective_mae": scaf_est["objective_value"],
                  "solver": scaf_est["solver"], **dict(zip(base, scaf_est["weights"]))}]
    if reported:
        vb = stack(ens, ("validation",))
        va, vp = fusion.block_values(val.Y, vb, resolution, energy_fn)
        w_rep = np.array([reported[m] for m in base], dtype=float)
        scaf_rows.append({"plant": plant.name, "target": target,
                          "coefficients": f"reported for target {f.scaf_reported_target or 'unconfirmed'}",
                          "used": use_reported, "resolution": resolution,
                          "validation_objective_mae": fusion.objective_value(w_rep / w_rep.sum(), va.reshape(-1),
                                                                             vp.reshape(-1, len(base)), "mae"),
                          "solver": "configured", **dict(zip(base, w_rep))})
    fused = {"UEBF": fusion.fuse(base_blocks, fusion.uniform_weights(len(base))),
             "SCAF": fusion.fuse(base_blocks, w_scaf), "DALF": dalf_fc}
    forecasts = {m: np.concatenate([ens[m]["validation"], ens[m]["test"]]) for m in cfg.model.models}
    forecasts.update(fused)
    test_fc = {name: fc[val.n:] for name, fc in forecasts.items()}
    val_fc = {name: fc[:val.n] for name, fc in forecasts.items()}
    for name, fc in test_fc.items():
        check(f"{name}: 24 finite hourly forecasts per test day", fc.shape == test.Y.shape and bool(np.all(np.isfinite(fc))))

    # prediction tables
    pred_dir = paths.predictions_for(plant.name, target)
    modes = dalf_w["window_mode"]
    test_frame = prediction_frame(plant.name, target, test, test_fc, unit, modes[partitions == "test"].reset_index(drop=True))
    val_frame = prediction_frame(plant.name, target, val, val_fc, unit, modes[partitions == "validation"].reset_index(drop=True))
    val_frame["note"] = "SCAF weights were estimated on this period (in-sample); DALF uses only earlier days"
    test_frame.to_csv(pred_dir / "test_predictions_hourly.csv", index=False)
    val_frame.to_csv(pred_dir / "validation_predictions_hourly.csv", index=False)
    daily_frame = None
    if energy_fn is not None:
        daily_frame = pd.DataFrame({"plant": plant.name, "forecast_day": pd.to_datetime(test.origin_times).normalize(),
                                    "forecast_origin": test.origin_times, "actual": energy_fn(test.Y)})
        for name, fc in test_fc.items():
            daily_frame[name] = energy_fn(fc)
        for name in FUSION_METHODS:
            daily_frame[f"residual_{name}"] = daily_frame["actual"] - daily_frame[name]
        daily_frame["unit"] = e_unit
        check("daily energy = sum of 24 hourly values x 1 h", bool(np.allclose(daily_frame["actual"], test.Y.sum(axis=1))),
              f"{cfg.data.power_measurement} in {cfg.data.power_unit} -> {e_unit}")
        daily_frame.to_csv(pred_dir / "test_daily_energy.csv", index=False)

    # metrics (test period)
    times = test.target_times()
    rows = []
    for name, fc in test_fc.items():
        ident = {"plant": plant.name, "target": target, "model": name, "family": family_of(name),
                 "prediction": ("rolling fusion of seed means" if name == "DALF" else
                                "fixed fusion of seed means" if name in FUSION_METHODS else
                                f"mean of {len(cfg.training.seeds)} seeds"), "partition": "test"}
        rows += ev.metric_rows(ident, test.Y, fc, unit, "physical", "hourly", times)
        rows += ev.metric_rows(ident, prep.scale_target(test.Y, target_scaler), prep.scale_target(fc, target_scaler),
                               "dimensionless", "normalized", "hourly", times)
        if energy_fn is not None:
            rows += ev.metric_rows(ident, energy_fn(test.Y), energy_fn(fc), e_unit, "physical", "daily_energy",
                                   test.origin_times)
    metrics = pd.DataFrame(rows)
    metrics["primary"] = metrics["resolution"] == resolution
    check("all compared forecasts share one test timestamp set",
          metrics.groupby(["resolution", "scale"])["timestamp_signature"].nunique().eq(1).all())

    # per-seed test metrics for the run records
    seed_rows = []
    for m in cfg.model.models:
        for seed, fc_seed, record in zip(cfg.training.seeds, loaded[m]["test"], loaded[m]["records"]):
            entry = {"plant": plant.name, "target": target, "model": m, "variant": "main", "seed": seed}
            entry.update({f"test_{k}_hourly_{unit}": v for k, v in ev.compute_metrics(test.Y, fc_seed).items()})
            if energy_fn is not None:
                entry.update({f"test_{k}_daily_{e_unit}": v
                              for k, v in ev.compute_metrics(energy_fn(test.Y), energy_fn(fc_seed)).items()})
            seed_rows.append(entry)

    # base model with the lowest validation MAE (selection never uses test data)
    val_mae = {}
    for m in base:
        a, p = fusion.block_values(val.Y, ens[m]["validation"][..., None], resolution, energy_fn)
        val_mae[m] = float(np.mean(np.abs(a.reshape(-1) - p.reshape(-1))))
    best_base = min(val_mae, key=val_mae.get)
    write_json(paths.tables_for(plant.name, target) / "base_model_selection.json",
               {"selected_on": "validation MAE at the fusion resolution", "validation_mae": val_mae,
                "best_base_model": best_base, "resolution": resolution})

    # statistical tests
    stats_rows = []
    if cfg.stats.enabled:
        for baseline in [b for b in cfg.stats.baselines if b in test_fc and b != cfg.stats.candidate]:
            for loss in cfg.stats.loss_functions:
                lb = st.daily_losses(test.Y, test_fc[baseline], resolution, loss, energy_fn)
                lc = st.daily_losses(test.Y, test_fc[cfg.stats.candidate], resolution, loss, energy_fn)
                result = st.diebold_mariano(lb, lc, horizon=1, lag=cfg.stats.hac_lag,
                                            harvey_correction=cfg.stats.harvey_correction)
                row = {"plant": plant.name, "target": target, "resolution": resolution, "loss": loss,
                       "candidate": cfg.stats.candidate, "baseline": baseline,
                       "loss_differential": "d_t = L(baseline) - L(candidate); positive = candidate better", **result}
                if cfg.stats.wilcoxon:
                    row.update(st.wilcoxon_paired(lb, lc))
                stats_rows.append(row)

    # fusion-level ablation 
    ablation = []

    def ablation_rows(experiment: str, setting, is_main: bool, fc_by_method: dict):
        for method, fc in fc_by_method.items():
            if resolution == "daily_energy":
                a, p, u, t = energy_fn(test.Y), energy_fn(fc), e_unit, test.origin_times
            else:
                a, p, u, t = test.Y, fc, unit, times
            for r in ev.metric_rows({"plant": plant.name, "target": target, "experiment": experiment,
                                     "setting": str(setting), "is_main_setting": is_main, "method": method},
                                    a, p, u, "physical", resolution, t):
                ablation.append(r)

    for wd in cfg.ablation.dalf_window_days:
        fc_wd, _ = run_dalf(base_blocks, window_days=wd)
        ablation_rows("dalf_window_days", wd, wd == f.window_days, {"DALF": fc_wd[val.n:]})
    if cfg.ablation.mse_dalf:
        fc_mse, w_mse = run_dalf(base_blocks, objective="mse")
        ablation_rows("dalf_objective (separate experiment)", "mse", False, {"DALF": fc_mse[val.n:]})
        ablation_rows("dalf_objective (separate experiment)", "mae", True, {"DALF": test_fc["DALF"]})
        w_mse.to_csv(paths.tables_for(plant.name, target) / "dalf_weights_mse_experiment.csv", index=False)
    ablation_rows("model_and_fusion_comparison", "main", True, test_fc)
    for frac in masking:
        if not all(frac in loaded[m]["masked"] for m in base):
            continue
        masked_test = {m: loaded[m]["masked"][frac].mean(axis=0) for m in base}
        masked_blocks = np.concatenate([np.stack([ens[m]["validation"] for m in base], axis=-1),
                                        np.stack([masked_test[m] for m in base], axis=-1)])
        fc_mask, _ = run_dalf(masked_blocks)
        mb = masked_blocks[val.n:]
        ablation_rows("input_masking_fraction", frac, False,
                      {**{m: masked_test[m] for m in base}, "UEBF": fusion.fuse(mb, fusion.uniform_weights(len(base))),
                       "SCAF": fusion.fuse(mb, w_scaf), "DALF": fc_mask[val.n:]})
    variant_rows = []
    variants = sorted({p.name.split("_")[1] for p in run_dirs["predictions"].glob("*_seed*.npz")} - {"main"})
    for variant in variants:
        for m in [x for x in cfg.model.models if x in HYBRID_MODELS]:
            seeds = cfg.ablation.variant_seeds
            try:
                var = load_seed_predictions(run_dirs, m, variant, seeds, sets, ())
            except FileNotFoundError:
                continue
            idx = [cfg.training.seeds.index(s) for s in seeds if s in cfg.training.seeds]
            for label, fc in ((variant, var["test"].mean(axis=0)),
                              ("main (same seeds)", loaded[m]["test"][idx].mean(axis=0) if idx else None)):
                if fc is None:
                    continue
                if resolution == "daily_energy":
                    a, p, u = energy_fn(test.Y), energy_fn(fc), e_unit
                else:
                    a, p, u = test.Y, fc, unit
                for k, v in ev.compute_metrics(a, p).items():
                    variant_rows.append({"plant": plant.name, "target": target, "model": m, "variant": label,
                                         "seeds": ",".join(map(str, seeds)), "resolution": resolution, "metric": k,
                                         "value": v, "unit": ev.metric_unit(k, u)})

    # computational cost
    records = [r for m in cfg.model.models for r in loaded[m]["records"]]
    test_dalf = dalf_w[dalf_w["partition"] == "test"]
    base_inference = sum(r["inference_seconds_test"] for m in base for r in loaded[m]["records"])
    dalf_cost = {"plant": plant.name, "target": target, "test_days": test.n,
                 "total_optimization_seconds": float(test_dalf["optimization_seconds"].sum()),
                 "mean_ms_per_day": 1000.0 * float(test_dalf["optimization_seconds"].mean()),
                 "values_per_optimization": int(test_dalf["values_in_objective"].median()),
                 "ensemble_forecast_seconds_test": base_inference + float(test_dalf["optimization_seconds"].sum()),
                 "ensemble_note": (f"inference of {len(base)} base models x {len(cfg.training.seeds)} seeds on all test "
                                   "days plus DALF optimization; single-model inference is in computational_cost_runs")}

    # per-target tables
    tdir = paths.tables_for(plant.name, target)
    metrics.to_csv(tdir / "metrics.csv", index=False)
    dalf_w.to_csv(tdir / "dalf_weights.csv", index=False)
    pd.DataFrame(scaf_rows).to_csv(tdir / "scaf_weights.csv", index=False)
    log.info("  %s/%s: SCAF %s; DALF %d test days (%s); best base model on validation: %s", plant.name, target,
             scaf_source, test.n, dalf_w.loc[dalf_w["partition"] == "test", "solver"].value_counts().to_dict(), best_base)
    summary = metrics[metrics["primary"] & (metrics["scale"] == "physical") & (metrics["metric"] == "MAE")]
    log.info("  test MAE at %s resolution:\n%s", resolution,
             summary[["model", "value", "unit"]].to_string(index=False))
    return {"metrics": metrics, "seed_rows": seed_rows, "records": records, "stats": stats_rows,
            "dalf_weights": dalf_w.assign(plant=plant.name, target=target), "scaf": scaf_rows,
            "ablation": ablation, "variants": variant_rows, "dalf_cost": dalf_cost, "checks": checks,
            "test_frame": test_frame, "daily_frame": daily_frame}


def architecture_table(cfg: ExperimentConfig, records: list[dict]) -> pd.DataFrame:
    rows = []
    for target in cfg.targets:
        n_features = len(prep.variables_for_target(target))
        for model in cfg.model.models:
            params = [r.get("trainable_parameters") for r in records
                      if r["model"] == model and r["target"] == target and r.get("variant") == "main"]
            for i, layer in enumerate(model_lib.architecture_spec(model, cfg.model, n_features, cfg.window.input_hours,
                                                                  cfg.window.horizon_hours)):
                rows.append({"target": target, "model": model, "family": model_lib.model_family(model), "position": i,
                             **layer, "trainable_parameters": params[0] if params else None,
                             "status": "IMPLEMENTATION DEFAULT - align with the manuscript"})
    return pd.DataFrame(rows)


def hyperparameter_table(cfg: ExperimentConfig) -> pd.DataFrame:
    rows = []
    as_dict = C.config_to_dict(cfg)
    for section in ("window", "model", "training", "fusion", "split", "cleaning", "stats", "ablation"):
        for key, value in as_dict[section].items():
            rows.append({"section": section, "setting": key,
                         "value": json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value})
    return pd.DataFrame(rows)


def stage_evaluate(cfg: ExperimentConfig, paths: Paths, log: logging.Logger, args) -> None:
    C.require_confirmed_site_settings(cfg)
    collected = {k: [] for k in ("metrics", "seed_rows", "records", "stats", "dalf_weights", "scaf", "ablation",
                                 "variants", "dalf_cost", "checks")}
    prediction_tables = {}
    for plant in cfg.plants:
        for target in cfg.targets:
            log.info("Evaluating %s/%s", plant.name, target)
            out = evaluate_target(cfg, paths, plant, target, log)
            for key in collected:
                value = out[key]
                collected[key].extend(value if isinstance(value, list) else [value])
            prediction_tables[f"{plant.name}_{target}_hourly"] = out["test_frame"]
            if out["daily_frame"] is not None:
                prediction_tables[f"{plant.name}_{target}_daily"] = out["daily_frame"]
    metrics = pd.concat(collected["metrics"], ignore_index=True)
    improvements = ev.improvement_table(metrics, "DALF")
    if len(improvements):
        primary = {t: cfg.fusion.resolution[t] for t in cfg.targets}
        improvements["primary"] = improvements.apply(lambda r: r["resolution"] == primary[r["target"]], axis=1)
    fusion_cmp = metrics[metrics["model"].isin(FUSION_METHODS)].pivot_table(
        index=["plant", "target", "resolution", "scale", "metric", "unit"], columns="model", values="value").reset_index()
    for b in ("SCAF", "UEBF"):
        if b in fusion_cmp and "DALF" in fusion_cmp:
            fusion_cmp[f"improvement_vs_{b}_pct"] = (fusion_cmp[b] - fusion_cmp["DALF"]) / fusion_cmp[b] * 100.0
    stats_table = st.apply_holm(pd.DataFrame(collected["stats"]), cfg.stats.alpha,
                                family_columns=("plant", "target", "resolution", "loss")) if collected["stats"] else pd.DataFrame()
    records = pd.DataFrame(collected["records"]).drop(columns=["model_config"], errors="ignore")
    run_table = records.merge(pd.DataFrame(collected["seed_rows"]), on=["plant", "target", "model", "variant", "seed"],
                              how="left") if len(records) else records
    cost_runs = records[["plant", "target", "model", "family", "variant", "seed", "trainable_parameters",
                         "total_parameters", "training_seconds", "epochs_run", "best_epoch", "model_file_bytes",
                         "inference_seconds_test", "inference_ms_per_test_window"]].copy() if len(records) else records
    cost_summary = cost_runs[cost_runs["variant"] == "main"].groupby(["model", "family"]).agg(
        trainable_parameters=("trainable_parameters", "first"),
        mean_training_seconds=("training_seconds", "mean"), std_training_seconds=("training_seconds", "std"),
        mean_model_file_kb=("model_file_bytes", lambda x: x.mean() / 1024),
        mean_inference_ms_per_window=("inference_ms_per_test_window", "mean"), runs=("seed", "size")).reset_index() \
        if len(cost_runs) else cost_runs
    tables = {
        "metrics_all": metrics, "improvements": improvements, "fusion_comparison": fusion_cmp,
        "statistical_tests": stats_table, "scaf_weights": pd.DataFrame(collected["scaf"]),
        "dalf_weights": pd.concat(collected["dalf_weights"], ignore_index=True),
        "ablation_fusion": pd.DataFrame(collected["ablation"]),
        "ablation_model_variants": pd.DataFrame(collected["variants"]),
        "run_records": run_table, "computational_cost_runs": cost_runs, "computational_cost_summary": cost_summary,
        "dalf_optimization_cost": pd.DataFrame(collected["dalf_cost"]),
        "model_architectures": architecture_table(cfg, collected["records"]),
        "hyperparameters": hyperparameter_table(cfg),
        "metric_definitions": pd.DataFrame([{"metric": k, "definition": v} for k, v in ev.METRIC_DEFINITIONS.items()]),
    }
    for name in ("data_quality_summary", "sample_counts", "ablation_interpolation_limits", "pipeline_runtime"):
        file = paths.tables / f"{name}.csv"
        if file.exists():
            tables[name] = pd.read_csv(file)
    for name, df in tables.items():
        if name not in ("data_quality_summary", "sample_counts", "ablation_interpolation_limits", "pipeline_runtime"):
            save_table(df, paths.tables / f"{name}.csv")
    write_json(paths.logs / "evaluation_checks.json", collected["checks"])
    write_excel(tables, paths.tables / "all_tables.xlsx", log)
    write_excel(prediction_tables, paths.tables / "prediction_tables.xlsx", log)
    log.info("Evaluation checks: %d passed", len(collected["checks"]))


#  plot
def history_frame(paths: Paths, plant: str, target: str) -> pd.DataFrame:
    run_dirs = paths.run_dirs(plant, target)
    frames = []
    for rec_file in sorted(run_dirs["histories"].glob("*_main_seed*_record.json")):
        rec = json.loads(rec_file.read_text(encoding="utf-8"))
        hist_file = run_dirs["histories"] / rec_file.name.replace("_record.json", "_history.csv")
        if hist_file.exists():
            h = pd.read_csv(hist_file)
            frames.append(h.assign(model=rec["model"], seed=rec["seed"]))
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def stage_plot(cfg: ExperimentConfig, paths: Paths, log: logging.Logger, args) -> None:
    viz.setup_style()
    saver = viz.FigureSaver(cfg.output.figure_formats, cfg.output.dpi, cfg.synthetic_data)
    for old in paths.figures.rglob("captions.md"):
        old.unlink()
    t = paths.tables
    required = ["metrics_all", "improvements", "model_architectures", "computational_cost_runs", "dalf_optimization_cost"]
    missing = [n for n in required if not (t / f"{n}.csv").exists()]
    if missing:
        raise FileNotFoundError(f"Tables {missing} not found in {t}; run --stage evaluate first.")
    metrics, improvements = pd.read_csv(t / "metrics_all.csv"), pd.read_csv(t / "improvements.csv")
    s = cfg.split
    viz.plot_workflow(saver, paths.figures, {"train": s.train_fraction, "validation": s.validation_fraction,
                                             "test": s.test_fraction, "input_hours": cfg.window.input_hours,
                                             "horizon": cfg.window.horizon_hours, "window_days": cfg.fusion.window_days})
    arch = pd.read_csv(t / "model_architectures.csv")
    viz.plot_architectures(saver, paths.figures, arch[arch["target"] == ("power" if "power" in cfg.targets else cfg.targets[0])])
    viz.plot_fusion_comparison(saver, paths.figures, metrics, improvements)
    viz.plot_error_heatmap(saver, paths.figures, metrics)
    cost = pd.read_csv(t / "computational_cost_runs.csv")
    viz.plot_computational_cost(saver, paths.figures, cost[cost["variant"] == "main"], pd.read_csv(t / "dalf_optimization_cost.csv"))
    for plant in cfg.plants:
        for target in cfg.targets:
            pdir, fdir, tdir = paths.predictions_for(plant.name, target), paths.figures_for(plant.name, target), \
                paths.tables_for(plant.name, target)
            hourly = pd.read_csv(pdir / "test_predictions_hourly.csv", parse_dates=["forecast_origin", "timestamp"])
            weights = pd.read_csv(tdir / "dalf_weights.csv", parse_dates=["forecast_origin"])
            best = json.loads((tdir / "base_model_selection.json").read_text(encoding="utf-8"))["best_base_model"]
            unit = C.target_unit(cfg, target)
            name = f"{plant.name} - {'PV power production' if target == 'power' else 'solar irradiance'}"
            methods = ["DALF", "SCAF", "UEBF", best]
            daily_file = pdir / "test_daily_energy.csv"
            if target == "power" and daily_file.exists():
                daily = pd.read_csv(daily_file, parse_dates=["forecast_day", "forecast_origin"])
                e_unit = daily["unit"].iloc[0]
                viz.plot_daily_energy(saver, fdir, daily, f"{name}: daily energy", e_unit)
                viz.plot_hourly_series(saver, fdir, hourly, "figS1_hourly_power_actual_vs_predicted",
                                       f"{name}: hourly power", f"Power ({unit})", cfg.output.zoom_days, "S1")
                err = daily[["forecast_day"]].copy()
                for m in FUSION_METHODS:
                    err[m] = (daily["actual"] - daily[m]).abs()
                viz.plot_error_timeseries(saver, fdir, err, f"{name}: absolute daily energy error",
                                          f"|actual - forecast| ({e_unit})")
                viz.plot_residuals(saver, fdir, daily, methods, f"{name}: daily energy residuals", e_unit, "days")
                viz.plot_scatter(saver, fdir, daily, methods, f"{name}: daily energy", e_unit, "days")
            else:
                label = "Solar irradiance" if target == "irradiance" else "Power"
                viz.plot_hourly_series(saver, fdir, hourly, "fig04_hourly_actual_vs_predicted", f"{name}: hourly values",
                                       f"{label} ({unit})", cfg.output.zoom_days, 4)
                err = hourly.assign(forecast_day=hourly["forecast_origin"].dt.normalize()).groupby("forecast_day")
                err = pd.DataFrame({m: err.apply(lambda g, m=m: float(np.mean(np.abs(g["actual"] - g[m]))))
                                    for m in FUSION_METHODS}).reset_index()
                viz.plot_error_timeseries(saver, fdir, err, f"{name}: daily mean absolute error", f"Daily MAE ({unit})")
                viz.plot_residuals(saver, fdir, hourly, methods, f"{name}: hourly residuals", unit, "hours")
                viz.plot_scatter(saver, fdir, hourly, methods, f"{name}: hourly values", unit, "hours")
            histories = history_frame(paths, plant.name, target)
            if len(histories):
                viz.plot_loss_curves(saver, fdir, histories, f"{name}: training and validation loss")
            viz.plot_weights(saver, fdir, weights, list(cfg.fusion.base_models), f"{name}: DALF weights")
    n_files = sum(1 for _ in paths.figures.rglob("*.png"))
    log.info("Figures written: %d PNG files (plus %s) under %s", n_files,
             ", ".join(x for x in cfg.output.figure_formats if x != "png") or "no vector copies", paths.figures)



STAGE_FUNCTIONS = {"preprocess": stage_preprocess, "train": stage_train, "evaluate": stage_evaluate, "plot": stage_plot}


def main(argv=None) -> int:
    args = parse_args(argv)
    if args.write_default_config:
        C.save_config(ExperimentConfig(), args.write_default_config)
        print(f"Default configuration written to {Path(args.write_default_config).resolve()}")
        return 0
    try:
        cfg = build_config(args)
    except (ConfigError, data_loader.DataValidationError) as exc:
        print(f"Configuration error:\n{exc}", file=sys.stderr)
        return 2
    paths = Paths(cfg.output.output_dir)
    log = setup_logging(paths.logs)
    if cfg.synthetic_data:
        log.warning("SMOKE TEST ON SYNTHETIC DATA: outputs check the code only and are NOT research results.")
    C.save_config(cfg, paths.logs / "config_used.json")
    write_json(paths.logs / "run_metadata.json", {**environment_info(), "command": sys.argv,
                                                  "started": pd.Timestamp.now().isoformat(timespec="seconds"),
                                                  "synthetic_data": cfg.synthetic_data})
    stages = STAGES if args.stage == "all" else (args.stage,)
    started = time.perf_counter()
    try:
        for stage in stages:
            log.info("========== stage: %s ==========", stage)
            t0 = time.perf_counter()
            STAGE_FUNCTIONS[stage](cfg, paths, log, args)
            record_stage(paths, stage, time.perf_counter() - t0)
            log.info("stage %s finished in %.1f s (peak process memory %s MB)", stage, time.perf_counter() - t0,
                     peak_memory_mb())
    except (ConfigError, data_loader.DataValidationError, FileNotFoundError) as exc:
        log.error("%s", exc)
        return 2
    except Exception:
        log.exception("The pipeline stopped with an error")
        return 1
    log.info("Completed %s in %.1f s. Outputs: %s", ", ".join(stages), time.perf_counter() - started,
             paths.root.resolve())
    return 0


if __name__ == "__main__":
    sys.exit(main())
